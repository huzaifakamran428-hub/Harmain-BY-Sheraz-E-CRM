"""
Excel (.xlsx) export via openpyxl. Kept in the infrastructure layer,
same reasoning as infrastructure/pdf_generator.py: the domain layer
must never import a third-party library directly.
"""
from __future__ import annotations

from pathlib import Path
from typing import Sequence

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from domain.models import Pilgrim
from core.exceptions import ExportError

_HEADER_FILL = PatternFill(start_color="1F2937", end_color="1F2937", fill_type="solid")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_TITLE_FONT = Font(bold=True, size=14)

# Column order/labels mirror the on-screen ledger table (LEDGER_COLUMNS /
# FREE_COLUMNS in presentation/widgets/pilgrim_table.py) minus the Photo
# column, which has no meaningful spreadsheet representation.
_COLUMNS = [
    ("SR", lambda p: (p.free_serial
                      if (getattr(p, "free_serial", 0)
                          and str(getattr(p.status, "value", p.status)) in ("FreeMale", "FreeFemale"))
                      else p.serial_number)),
    ("Date", lambda p: p.permit_date.isoformat() if p.permit_date else ""),
    ("Time", lambda p: p.permit_time.strftime("%H:%M") if p.permit_time else ""),
    ("Batch", lambda p: p.batch),
    ("Name", lambda p: p.name),
    ("Gender", lambda p: p.gender.value if hasattr(p.gender, "value") else p.gender),
    ("Passport No.", lambda p: p.passport_number),
    ("Visa No.", lambda p: p.visa_number),
    ("Email", lambda p: p.email),
    ("WhatsApp", lambda p: p.whatsapp_number),
    ("Charge", lambda p: p.charge),
    ("Received", lambda p: p.total_received),
    ("Remaining", lambda p: p.remaining),
    ("Status", lambda p: p.status.value if hasattr(p.status, "value") else p.status),
]


def _write_sheet(ws, title: str, pilgrims: Sequence[Pilgrim]) -> None:
    ws.title = title[:31]  # Excel sheet-name length limit

    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(_COLUMNS))
    title_cell = ws.cell(row=1, column=1, value=title)
    title_cell.font = _TITLE_FONT

    header_row = 3
    for col_index, (label, _) in enumerate(_COLUMNS, start=1):
        cell = ws.cell(row=header_row, column=col_index, value=label)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = Alignment(horizontal="center")

    for row_offset, pilgrim in enumerate(pilgrims, start=header_row + 1):
        for col_index, (_, getter) in enumerate(_COLUMNS, start=1):
            ws.cell(row=row_offset, column=col_index, value=getter(pilgrim))

    last_row = header_row + len(pilgrims) + 1
    ws.cell(row=last_row, column=1, value="Total").font = Font(bold=True)
    ws.cell(row=last_row, column=len(_COLUMNS), value=len(pilgrims)).font = Font(bold=True)

    for col_index, (label, _) in enumerate(_COLUMNS, start=1):
        width = max(len(label) + 4, 12)
        ws.column_dimensions[get_column_letter(col_index)].width = width
    ws.freeze_panes = ws.cell(row=header_row + 1, column=1)


def export_party_excel(party_name: str, pilgrims: Sequence[Pilgrim], output_path: Path) -> Path:
    """One sheet, titled with the party's name, listing every pilgrim on its ledger."""
    try:
        wb = Workbook()
        _write_sheet(wb.active, f"{party_name} - Pilgrims", pilgrims)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(output_path))
        return output_path
    except OSError as e:
        raise ExportError(f"Could not save the Excel file: {e}") from e


def export_free_account_excel(male: Sequence[Pilgrim], female: Sequence[Pilgrim],
                               output_path: Path) -> Path:
    """Two sheets - Male and Female - each with its own list, as shown in the app."""
    try:
        wb = Workbook()
        _write_sheet(wb.active, "Male", male)
        female_sheet = wb.create_sheet()
        _write_sheet(female_sheet, "Female", female)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        wb.save(str(output_path))
        return output_path
    except OSError as e:
        raise ExportError(f"Could not save the Excel file: {e}") from e
